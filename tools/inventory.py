"""Read-only inventory of the immutable source workbooks."""
from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
for path in sorted((ROOT / "sources" / "original").glob("*.xlsx")):
    wb = load_workbook(path, data_only=False, read_only=False)
    print(f"WORKBOOK {path.name}")
    print(f"SHEETS {len(wb.worksheets)} DEFINED_NAMES {len(wb.defined_names)}")
    for name, value in wb.defined_names.items():
        print("NAME", name, value.attr_text)
    for ws in wb.worksheets:
        print(f"SHEET {ws.title!r} state={ws.sheet_state} rows={ws.max_row} cols={ws.max_column} merged={len(ws.merged_cells.ranges)} validations={len(ws.data_validations.dataValidation) if ws.data_validations else 0} images={len(ws._images)} charts={len(ws._charts)}")
        cells = [(c.coordinate, str(c.value)[:170]) for row in ws for c in row if c.value is not None]
        print("NONEMPTY", len(cells), "FORMULAS", sum(v.startswith('=') for _,v in cells), "COMMENTS", sum(c.comment is not None for row in ws for c in row))
        print("SAMPLE", json.dumps(cells[:80], ensure_ascii=False))
        for dv in ws.data_validations.dataValidation if ws.data_validations else []:
            print("VALIDATION", str(dv.sqref), dv.type, dv.formula1, dv.formula2)

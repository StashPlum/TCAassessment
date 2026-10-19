"""Generate disposable, fictional stage-three export examples for visual QA."""
import os
import sys
from io import BytesIO
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
os.environ.setdefault('DJANGO_SETTINGS_MODULE','config.settings')
import django
django.setup()
from openpyxl import load_workbook
from pypdf import PdfReader
from core.assessment_documents import assessment_records, pdf_bytes, preparation_records, xlsx_bytes
from core.models import Assessment, Site
from core.preparation_service import preparation_items

site=Site.objects.get(code='fictional-stage-one')
items=preparation_items(site,{},selected=['1.C.1','15.A.10'])
records=preparation_records(site,items)
title=f'{site.organization} | TCA evidence preparation | {site.name}'
scope='Selected requirements across pages (2)'
folder=Path('output/stage3-qa');folder.mkdir(parents=True,exist_ok=True)
pdf=pdf_bytes(title,records,scope)
xlsx=xlsx_bytes(title,records,scope)
(folder/'preparation-example.pdf').write_bytes(pdf)
(folder/'preparation-example.xlsx').write_bytes(xlsx)
sheet=load_workbook(BytesIO(xlsx),read_only=True,data_only=True).active
assert sheet['B3'].value==2 and {sheet['F6'].value,sheet['F7'].value}=={'1.C.1','15.A.10'}
text='\n'.join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages)
assert 'CEVA Logistics' in text and '1.C.1' in text and '15.A.10' in text
print(f'Generated {len(records)} complete fictional records; PDF pages={len(PdfReader(BytesIO(pdf)).pages)}; XLSX rows={sheet.max_row}')
approved=Assessment.objects.filter(site=site,status='approved').order_by('-version').first()
if approved:
    archive=assessment_records(approved)
    archive_pdf=pdf_bytes('CEVA Logistics | TCA assessment history',archive,
                          f'{site.name} v{approved.version} ({approved.status}, {approved.policy_version})')
    (folder/'history-example.pdf').write_bytes(archive_pdf)
    archive_xlsx=xlsx_bytes('CEVA Logistics | TCA assessment history',archive,
                             f'{site.name} v{approved.version} ({approved.status}, {approved.policy_version})')
    (folder/'history-example.xlsx').write_bytes(archive_xlsx)
    archive_sheet=load_workbook(BytesIO(archive_xlsx),read_only=True,data_only=True).active
    assert archive_sheet['B3'].value==len(archive)==150
    assert 'advisory-v2' in '\n'.join(page.extract_text() for page in PdfReader(BytesIO(archive_pdf)).pages)
    print(f'Historical approved v{approved.version}: PDF pages={len(PdfReader(BytesIO(archive_pdf)).pages)}; XLSX records={archive_sheet["B3"].value}')

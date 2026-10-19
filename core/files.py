"""Private immutable bytes and deliberate CSV/XLSX preview, no formula execution."""
import csv,hashlib,zipfile
from io import BytesIO,StringIO
from pathlib import Path
from django import forms
from PIL import Image

def checked_upload(upload):
    if upload.size>10*1024*1024: raise forms.ValidationError('Local pilot limit: 10 MB per file.')
    raw=upload.read();upload.seek(0);suffix=Path(upload.name).suffix.lower()
    if suffix not in ('.pdf','.png','.jpg','.jpeg','.xlsx','.csv'): raise forms.ValidationError('Use PDF, PNG, JPG, XLSX or UTF-8 CSV.')
    if suffix=='.pdf' and not raw.startswith(b'%PDF-'): raise forms.ValidationError('The file content is not a PDF.')
    if suffix in ('.png','.jpg','.jpeg'):
        try:
            with Image.open(BytesIO(raw)) as img:
                if img.format not in ('PNG','JPEG') or img.width*img.height>25_000_000: raise ValueError()
                img.verify()
        except Exception: raise forms.ValidationError('The image is invalid or exceeds 25 megapixels.')
    if suffix=='.xlsx':
        try:
            with zipfile.ZipFile(BytesIO(raw)) as z:
                if sum(i.file_size for i in z.infolist())>40*1024*1024 or len(z.infolist())>1000: raise ValueError()
                if '[Content_Types].xml' not in z.namelist() or 'xl/workbook.xml' not in z.namelist(): raise ValueError()
                if any('vbaProject' in name for name in z.namelist()): raise ValueError()
        except Exception: raise forms.ValidationError('Invalid, macro-enabled or oversized expanded XLSX.')
    if suffix=='.csv':
        try: value=raw.decode('utf-8-sig')
        except UnicodeDecodeError: raise forms.ValidationError('CSV must use UTF-8 encoding.')
        if '\x00' in value: raise forms.ValidationError('CSV contains binary data.')
    return upload,hashlib.sha256(raw).hexdigest()

def read_table(attachment,sheet=''):
    with attachment.file.open('rb') as source: raw=source.read()
    if hashlib.sha256(raw).hexdigest()!=attachment.sha256: raise ValueError('Original file hash no longer matches; import stopped.')
    extension=Path(attachment.original_name).suffix.lower()
    if extension=='.csv':
        data=raw.decode('utf-8-sig')
        try: dialect=csv.Sniffer().sniff(data[:8192],delimiters=',;\t')
        except csv.Error: dialect=csv.excel
        values=list(csv.reader(StringIO(data),dialect));names=['CSV']
    elif extension=='.xlsx':
        from openpyxl import load_workbook
        workbook=load_workbook(BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        names=workbook.sheetnames
        if sheet and sheet not in names: workbook.close();raise ValueError('Choose an existing worksheet.')
        worksheet=workbook[sheet or names[0]]
        if worksheet.max_row>10001 or worksheet.max_column>50: workbook.close();raise ValueError('Import limit is 10,000 rows and 50 columns; select a smaller source file.')
        values=[[str(v) if v is not None else '' for v in row] for row in worksheet.iter_rows(values_only=True)];workbook.close()
    else: raise ValueError('Only CSV/XLSX can be previewed for import. Other attachments remain evidence only.')
    if not values or not values[0] or len(values)>10001 or max(map(len,values))>50: raise ValueError('Empty or oversized table. Maximum 10,000 rows / 50 columns.')
    headers=values[0];width=len(headers)
    if any(len(row)>width for row in values[1:]): raise ValueError('Rows contain more columns than the header; fix the source layout.')
    rows=[row+['']*(width-len(row)) for row in values[1:] if any(str(v).strip() for v in row)]
    return {'headers':headers,'rows':rows,'sheet':sheet or names[0],'sheets':names,'origin_sha256':attachment.sha256,'formula_policy':'Literal cells only; formulas are not evaluated.'}

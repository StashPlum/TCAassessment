"""Stage 4: Kaizen start, team/tasks, Gantt parity, A3/A4 rules, sensei review and downloads."""
import json, re, zipfile
from datetime import date, timedelta
from io import BytesIO
from uuid import uuid4
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from .models import (A3Report, Assessment, Kaizen, KaizenAction, Membership, Observation, Participant, Requirement, Response,
                     ReviewNote, Site, Worksheet)
from .kaizen import gantt, gate_list, review_checks
from .methodology import rules_for
from .toolbox_registry import REGISTRY, initial_payload
from .toolbox_visuals import render_svg

READY={'overall_problem':'Synthetic delay','specific_problem':'Synthetic handover','solution':'Synthetic standard','baseline':'10','actual':'4',
       'unit':'min','population':'Synthetic shifts','data_source':'Synthetic register','sustain':'Synthetic weekly check'}

@override_settings(MEDIA_ROOT='tmp/test-stage4-media')
class StageFourTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('import_catalog',verbosity=0)
        User=get_user_model()
        cls.site=Site.objects.create(code='fictional-s4',name='Fictional Stage Four Site',country_code='NL')
        cls.other=Site.objects.create(code='fictional-s4-other',name='Fictional Other Site',country_code='BE')
        cls.lead=User.objects.create_user('s4-lead',first_name='Mila',last_name='Janssen')
        cls.member=User.objects.create_user('s4-member',first_name='Daan',last_name='Visser')
        cls.outside_staff=User.objects.create_user('s4-not-on-team')
        cls.sensei=User.objects.create_user('s4-sensei',first_name='Ingrid',last_name='Holm')
        cls.foreign=User.objects.create_user('s4-foreign')
        for user in (cls.lead,cls.member,cls.outside_staff): Membership.objects.create(user=user,site=cls.site,role='staff')
        Membership.objects.create(user=cls.sensei,site=cls.site,role='coach')
        Membership.objects.create(user=cls.foreign,site=cls.other,role='staff')
        Membership.objects.create(user=cls.foreign,site=cls.other,role='coach')

    def create(self,profile='A4',ready=False,**extra):
        self.client.force_login(self.lead)
        response=self.client.post(reverse('kaizens',args=[self.site.code]),{'title':'Fictional stage four event '+str(uuid4())[:6],'profile':profile,**extra})
        self.assertEqual(response.status_code,302,getattr(response,'content',b'')[:400])
        event=Kaizen.objects.latest('pk')
        if ready:
            event.data={**event.data,**READY,'root_cause':'Synthetic investigated cause'}
            event.status='results_under_review';event.submitted_at=timezone.now();event.save()
            KaizenAction.objects.create(event=event,phase='DO',description='Synthetic do',owner='Mila Janssen')
            KaizenAction.objects.create(event=event,phase='ACT',description='Synthetic act',owner='Mila Janssen')
            for phase,value in [('before',10),('after',4)]:
                Observation.objects.create(event=event,phase=phase,measure='wait',value=value,unit='min',population='Synthetic shifts',period_start=date(2026,9,1),period_end=date(2026,9,2),data_source='Synthetic register')
        return event

    # W21 -----------------------------------------------------------------
    def test_title_only_start_suggests_input_view_and_links_requirement(self):
        self.client.force_login(self.lead)
        page=self.client.get(reverse('kaizens',args=[self.site.code]),{'requirement':'4.C.3'})
        self.assertContains(page,'Help me choose');self.assertContains(page,'TCA 4.C.3');self.assertContains(page,'/ 75')
        response=self.client.post(reverse('kaizens',args=[self.site.code]),{'title':'Fictional only a title','requirement':'4.C.3',
            'help_scope':'one','help_recurrence':'recurring','help_cause':'known'})
        event=Kaizen.objects.get(title='Fictional only a title')
        self.assertRedirects(response,reverse('kaizen_detail',args=[self.site.code,event.pk])+'?created=1',fetch_redirect_response=False)
        self.assertEqual((event.input_profile,event.profile),('medium','A4'))
        self.assertEqual(event.data['overall_problem'],'Fictional only a title')
        self.assertEqual(list(event.linked_requirements.values_list('code',flat=True)),['4.C.3'])
        lead=event.participants.get();self.assertEqual((lead.user,lead.kaizen_role),(self.lead,'Kaizen lead'))
        self.assertContains(self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk])+'?created=1'),'Draft saved')
        # An explicit choice wins over the suggestion; no Small=A4 mapping is applied.
        self.client.post(reverse('kaizens',args=[self.site.code]),{'title':'Fictional explicit small A3','input_profile':'small','profile':'A3','help_scope':'site'})
        self.assertEqual(Kaizen.objects.get(title='Fictional explicit small A3').profile,'A3')
        self.assertContains(self.client.post(reverse('kaizens',args=[self.site.code]),{'title':''}),'This field is required')

    # W24 / W25 -----------------------------------------------------------
    def test_team_membership_rights_and_my_tasks(self):
        event=self.create()
        add=reverse('kaizen_add',args=[self.site.code,event.pk,'member'])
        # A user from another site cannot be added, even with a forged identifier.
        self.client.post(add,{'event_revision':event.revision,'user':self.foreign.pk,'kaizen_role':'Team member'})
        self.assertFalse(event.participants.filter(user=self.foreign).exists())
        self.assertEqual(self.client.post(add,{'event_revision':event.revision,'user':self.member.pk,'kaizen_role':'Team member'}).status_code,302)
        event.refresh_from_db();self.assertTrue(event.participants.filter(user=self.member).exists())
        # Only the team edits a team-managed event; other site staff can read it.
        self.client.force_login(self.outside_staff)
        self.assertEqual(self.client.post(reverse('kaizen_autosave',args=[self.site.code,event.pk]),{'revision':event.revision,'solution':'Not on the team'}).status_code,403)
        self.assertEqual(self.client.post(reverse('kaizen_add',args=[self.site.code,event.pk,'action']),{'event_revision':event.revision}).status_code,404)
        self.assertEqual(self.client.post(add,{'event_revision':event.revision,'user':self.outside_staff.pk,'kaizen_role':'Team member'}).status_code,404)
        self.assertContains(self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk])),'Read-only access')
        # The lead assigns a task; it appears in the member's home task list.
        self.client.force_login(self.lead)
        team_member=event.participants.get(user=self.member)
        response=self.client.post(reverse('kaizen_add',args=[self.site.code,event.pk,'action']),{'event_revision':event.revision,'phase':'DO','description':'Fictional place agenda cards',
            'assignee':team_member.pk,'start_date':'2026-10-20','end_date':'2026-10-24','progress':'0','workspace_stage':'Do'})
        self.assertEqual(response.status_code,302)
        task=KaizenAction.objects.get(description='Fictional place agenda cards')
        self.assertEqual((task.assignee,task.owner),(team_member,'Daan Visser'))
        self.client.force_login(self.member)
        self.assertContains(self.client.get(reverse('home')),'Fictional place agenda cards')
        event.refresh_from_db()
        progress=reverse('task_progress',args=[self.site.code,event.pk,task.pk])
        self.assertEqual(self.client.post(progress,{'event_revision':event.revision,'progress':'60','next':'home'}).status_code,302)
        task.refresh_from_db();self.assertEqual(task.progress,60)
        # A stale revision never overwrites the newer value.
        self.client.post(progress,{'event_revision':event.revision,'progress':'20'})
        task.refresh_from_db();self.assertEqual(task.progress,60)
        self.assertEqual(self.client.post(progress,{'event_revision':event.revision+1,'progress':'101'}).status_code,422)
        # Another site's user is denied outright.
        self.client.force_login(self.foreign)
        self.assertEqual(self.client.post(progress,{'event_revision':event.revision+1,'progress':'80'}).status_code,404)

    def test_task_needs_owner_and_valid_dates(self):
        event=self.create()
        url=reverse('kaizen_add',args=[self.site.code,event.pk,'action'])
        response=self.client.post(url,{'event_revision':event.revision,'phase':'DO','description':'No owner'})
        self.assertContains(response,'Choose an owner from the team or name another owner.')
        lead=event.participants.get()
        response=self.client.post(url,{'event_revision':event.revision,'phase':'DO','description':'Bad dates','assignee':lead.pk,'start_date':'2026-10-24','end_date':'2026-10-20'})
        self.assertContains(response,'End date must be on or after start date.')
        self.assertFalse(event.actions.exists())

    # W26 -----------------------------------------------------------------
    def test_gantt_and_table_read_the_same_action_records(self):
        event=self.create()
        lead=event.participants.get();today=date(2026,10,19)
        a=KaizenAction.objects.create(event=event,phase='DO',description='First',owner='Mila Janssen',assignee=lead,start_date=today-timedelta(days=5),end_date=today+timedelta(days=4),progress=50)
        b=KaizenAction.objects.create(event=event,phase='ACT',description='Undated',owner='Mila Janssen')
        chart=gantt(list(event.actions.order_by('pk')),today)
        self.assertEqual([row['action'].pk for row in chart['rows']],[a.pk,b.pk])
        self.assertIsNone(chart['rows'][1]['bar'])
        span=(chart['end']-chart['start']).days+1
        self.assertAlmostEqual(chart['rows'][0]['bar']['width'],10/span*100,places=2)
        self.assertAlmostEqual(chart['rows'][0]['bar']['fill'],chart['rows'][0]['bar']['width']/2,places=2)
        table=self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk]),{'stage':'Do'})
        chart_page=self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk]),{'stage':'Do','view':'gantt'})
        for page in (table,chart_page):
            self.assertContains(page,'First');self.assertContains(page,'Undated')
        self.assertContains(chart_page,'Dates not set');self.assertContains(chart_page,'Table and Gantt show the same action records')
        # Editing through the table changes what the Gantt shows; there is no second plan store.
        event.refresh_from_db()
        self.client.post(reverse('kaizen_action_edit',args=[self.site.code,event.pk,a.pk]),{'event_revision':event.revision,'phase':'DO','description':'First renamed','assignee':lead.pk,'start_date':'2026-10-14','end_date':'2026-10-23','progress':'100'})
        self.assertContains(self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk]),{'view':'gantt'}),'First renamed')

    # W28 -----------------------------------------------------------------
    def test_a3_and_bp_require_cause_analysis_but_a4_does_not(self):
        self.assertEqual(rules_for('BP')['upload_type'],'A3 Kaizen');self.assertEqual(rules_for('A4')['cause_analysis'],'optional')
        a4=self.create('A4');a3=self.create('A3');bp=self.create('BP')
        for event in (a4,a3,bp): event.data.update(specific_problem='Synthetic');event.save()
        self.assertNotIn('root_cause',[g['key'] for g in gate_list(a4)])
        for event in (a3,bp): self.assertIn('Cause analysis (required for A3 and Business Process)',review_checks(event))
        # Plan to Do: A3 needs the cause analysis, A4 does not.
        for event,expected in [(a4,'planned'),(a3,'proposed')]:
            self.client.post(reverse('kaizen_detail',args=[self.site.code,event.pk]),{'revision':event.revision,'action':'advance'})
            event.refresh_from_db();self.assertEqual(event.status,expected)

    def test_view_and_type_switches_keep_every_record(self):
        event=self.create('A4')
        event.data.update(extra_detail='Keep me',root_cause='');event.save()
        lead=event.participants.get()
        KaizenAction.objects.create(event=event,phase='DO',description='Keep action',owner='Mila Janssen',assignee=lead,progress=40)
        Worksheet.objects.create(event=event,method='pareto',stage='Check',title='Keep worksheet',payload=initial_payload('pareto'))
        url=reverse('input_profile',args=[self.site.code,event.pk])
        self.assertEqual(self.client.post(url,{'revision':event.revision,'input_profile':'extensive','profile':'A3'}).status_code,302)
        event.refresh_from_db()
        self.assertEqual((event.input_profile,event.profile,event.data['extra_detail']),('extensive','A3','Keep me'))
        self.assertEqual(event.actions.get().progress,40);self.assertTrue(event.worksheets.exists());self.assertEqual(event.participants.count(),1)
        self.assertIn('root_cause',[g['key'] for g in gate_list(event)])
        self.client.post(url,{'revision':event.revision,'input_profile':'small','profile':'A4'})
        event.refresh_from_db();self.assertEqual((event.input_profile,event.profile),('small','A4'))
        self.assertEqual(self.client.post(url,{'revision':0,'input_profile':'medium'}).status_code,409)

    # W31-W33 --------------------------------------------------------------
    def test_queue_recent_first_status_filters_and_selected_export(self):
        old=self.create(ready=True);new=self.create(ready=True)
        Kaizen.objects.filter(pk=old.pk).update(submitted_at=timezone.now()-timedelta(days=5))
        Kaizen.objects.filter(pk=new.pk).update(submitted_at=timezone.now()-timedelta(hours=1))
        self.client.force_login(self.sensei)
        page=self.client.get(reverse('review_queue',args=[self.site.code]))
        ids=[row['event'].pk for row in page.context['rows']]
        self.assertEqual(ids[:2],[new.pk,old.pk])
        self.assertContains(page,'Submitted for review')
        export=reverse('review_export',args=[self.site.code,'xlsx'])
        workbook=load_workbook(BytesIO(self.client.get(export,{'ids':[new.pk,old.pk]}).content))
        self.assertEqual(workbook.sheetnames[:4],['Events','A3 summary','Actions','Gates'])
        rows=list(workbook['Events'].iter_rows(min_row=2,values_only=True))
        self.assertEqual({r[0] for r in rows},{f'KZ-{new.pk}',f'KZ-{old.pk}'})
        self.assertTrue(all(r[8]=='Submitted for review' for r in rows))
        pdf=self.client.get(reverse('review_export',args=[self.site.code,'pdf']),{'ids':[new.pk],'paper':'A4'})
        self.assertEqual(pdf['Content-Type'],'application/pdf')
        from pypdf import PdfReader
        reader=PdfReader(BytesIO(pdf.content));self.assertAlmostEqual(float(reader.pages[0].mediabox.width),841.89,places=1)
        self.assertIn('WORKING COPY',reader.pages[0].extract_text())
        # Another site's event identifier is refused.
        foreign=Kaizen.objects.create(site=self.other,title='Fictional foreign',profile='A4',created_by=self.foreign)
        self.assertEqual(self.client.get(export,{'ids':[new.pk,foreign.pk]}).status_code,404)
        # Returned state appears in its own filter.
        ReviewNote.objects.create(event=old,stage='Plan',kind='return',text='Synthetic return',actor=self.sensei,event_revision=0)
        Kaizen.objects.filter(pk=old.pk).update(status='in_progress')
        returned=self.client.get(reverse('review_queue',args=[self.site.code]),{'state':'returned'})
        self.assertEqual([row['event'].pk for row in returned.context['rows']],[old.pk])

    def test_return_needs_feedback_and_finalise_needs_every_gate(self):
        event=self.create(ready=True)
        approved=Assessment.objects.create(site=self.site,version=1,status='approved',created_by=self.sensei)
        Response.objects.create(assessment=approved,requirement=Requirement.objects.get(pk='4.C.3'),code='N',updated_by=self.sensei)
        event.linked_requirements.add('4.C.3')
        self.client.force_login(self.sensei)
        url=reverse('kaizen_review_detail',args=[self.site.code,event.pk])
        page=self.client.get(url);self.assertContains(page,'Gate checklist');self.assertContains(page,'Before and after')
        self.assertContains(self.client.post(url,{'revision':event.revision,'decision':'return'}),'Write at least one comment')
        # Finalising without the overall review fails and stores nothing.
        response=self.client.post(url,{'revision':event.revision,'decision':'finalize','feedback_Plan':'Kept text'})
        self.assertContains(response,'Finalising needs every gate');self.assertContains(response,'Kept text')
        self.assertFalse(A3Report.objects.filter(event=event).exists());self.assertFalse(event.review_notes.exists())
        # Return with feedback per phase: the team sees it.
        self.client.post(url,{'revision':event.revision,'decision':'return','feedback_Check':'Add five more after points'})
        event.refresh_from_db();self.assertEqual(event.status,'in_progress')
        self.client.force_login(self.lead)
        self.assertContains(self.client.get(reverse('kaizen_detail',args=[self.site.code,event.pk])),'Add five more after points')
        # Team resubmits; the sensei resolves the request and finalises.
        self.client.post(reverse('kaizen_detail',args=[self.site.code,event.pk]),{'revision':event.revision,'action':'submit_review'})
        event.refresh_from_db();self.assertEqual(event.status,'results_under_review')
        self.client.force_login(self.sensei)
        note=event.review_notes.get(kind='return')
        self.client.post(reverse('review_note',args=[self.site.code,event.pk]),{'revision':event.revision,'resolve':note.pk,'next':'review'})
        event.refresh_from_db()
        response=self.client.post(url,{'revision':event.revision,'decision':'finalize','coach_review':'Synthetic sensei conclusion'})
        self.assertEqual(response.status_code,302)
        event.refresh_from_db();self.assertEqual(event.status,'verified')
        self.assertEqual(sorted(event.a3_reports.values_list('paper',flat=True)),['A3','A4'])
        # Finalisation never changes a TCA answer.
        self.assertEqual(Response.objects.get(assessment=approved).code,'N')

    # W29 / W30 ------------------------------------------------------------
    def test_downloads_match_event_and_respect_site_access(self):
        event=self.create(ready=True)
        lead=event.participants.get()
        KaizenAction.objects.filter(event=event,phase='DO').update(assignee=lead,progress=80)
        pdf=reverse('event_export',args=[self.site.code,event.pk,'pdf']);xlsx=reverse('event_export',args=[self.site.code,event.pk,'xlsx'])
        from pypdf import PdfReader
        for paper,width in [('A3',1190.55),('A4',841.89)]:
            response=self.client.get(pdf,{'paper':paper});self.assertEqual(response.status_code,200)
            reader=PdfReader(BytesIO(response.content));text=''.join(p.extract_text() for p in reader.pages)
            self.assertAlmostEqual(float(reader.pages[0].mediabox.width),width,places=1)
            for value in ('WORKING COPY','Synthetic handover','Synthetic do','Complete content'): self.assertIn(value,text)
        self.assertEqual(self.client.get(pdf,{'paper':'A5'}).status_code,404)
        workbook=load_workbook(BytesIO(self.client.get(xlsx).content))
        actions=list(workbook['Actions'].iter_rows(min_row=2,values_only=True))
        self.assertEqual({r[2] for r in actions},{'Synthetic do','Synthetic act'})
        self.assertIn(80,[r[7] for r in actions])
        gates=list(workbook['Gates'].iter_rows(min_row=2,values_only=True))
        self.assertEqual(len(gates),len(gate_list(event)))
        summary={r[2]:r[3] for r in workbook['A3 summary'].iter_rows(min_row=2,values_only=True)}
        self.assertIn('Synthetic handover',summary['Problem'])
        # Full archive keeps team, assignment and progress.
        with zipfile.ZipFile(BytesIO(self.client.get(reverse('event_archive_zip',args=[self.site.code,event.pk])).content)) as archive:
            complete=json.loads(archive.read('event-complete.json'))
        self.assertTrue(any(a['progress']==80 and a['assignee_id']==lead.pk for a in complete['actions']))
        # Other-site users cannot download anything.
        self.client.force_login(self.foreign)
        for url,params in [(pdf,{'paper':'A3'}),(xlsx,{}),(reverse('kaizen_review_detail',args=[self.site.code,event.pk]),{}),(reverse('review_export',args=[self.site.code,'xlsx']),{'ids':[event.pk]})]:
            self.assertEqual(self.client.get(url,params).status_code,404)
        # Team members cannot open the sensei review page.
        self.client.force_login(self.lead)
        self.assertEqual(self.client.get(reverse('kaizen_review_detail',args=[self.site.code,event.pk])).status_code,404)

    # W22 / W23 ------------------------------------------------------------
    def test_every_method_has_help_and_long_labels_stay_inside_diagrams(self):
        self.client.force_login(self.lead)
        library=self.client.get(reverse('toolbox_library',args=[self.site.code]))
        self.assertEqual(len(library.context['specs']),45)
        for spec in REGISTRY.values():
            self.assertTrue(all(spec.get(k) for k in ('purpose','steps','pitfall','example','fields')),spec['key'])
        page=self.client.get(reverse('toolbox_library',args=[self.site.code]),{'method':'pareto'})
        for text in ('Pitfall.','Logistics example','Fields you fill in'): self.assertContains(page,text)
        self.assertEqual(len(self.client.get(reverse('toolbox_library',args=[self.site.code]),{'stage':'Check'}).context['specs']),library.context['counts']['Check'])
        long='Fictional very long label for a loading dock observation that keeps going for legibility checks'
        cases={'scatter':[{'x':i,'y':i*2,'x_unit':'n','y_unit':'min','period':f'{long} {i}','source':'Synthetic'} for i in range(1,4)],
               'boxplot':[{'group':long,'value':v,'unit':'min','period':'P','source':'S'} for v in (1,2,3,4,9)],
               'spaghetti':[{'route':'Picker','state':'current','point':f'{long} {i}','x':i,'y':i,'unit':'m','source':'S'} for i in range(3)]}
        for key,rows in cases.items():
            payload=initial_payload(key);payload['rows']=[{'id':str(uuid4()),'archived':False,**r} for r in rows]
            svg=render_svg(key,payload);width=float(re.search(r'viewBox="0 0 ([\d.]+)',svg).group(1))
            for x,size,value in re.findall(r'<text x="([\d.]+)" y="[-\d.]+" font-family="[^"]+" font-size="(\d+)" fill="[^"]+">([^<]*)</text>',svg):
                self.assertLessEqual(float(x)+len(value)*int(size)*.55,width+2,(key,value))
            self.assertIn('keeps going',svg,key)
        event=self.create()
        worksheet=Worksheet.objects.create(event=event,method='5whys',stage='Plan',title='5 Whys',payload=initial_payload('5whys'))
        self.assertContains(self.client.get(reverse('worksheet',args=[self.site.code,event.pk,worksheet.pk])),'Method help')

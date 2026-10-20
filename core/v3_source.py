"""Read and reconcile the TCA V3 assessment workbook without changing its wording.

The workbook is private and never committed. It is read read-only (never saved through openpyxl), or from a
JSON extract of its raw rows that carries the workbook's SHA-256. A numbered Master Order row starts a source
record; every following row up to the next numbered row is a continuation of that record's columns I/J/K.
"""
import hashlib
import json
import re
from collections import Counter, defaultdict
from pathlib import Path

SHEET = 'Blad1'
COLUMNS = ('Master Order', 'QuestionID', 'order', 'section', 'level', 'language', 'TCA Groups', 'Product Line',
           'v3 requirement', 'v3 evidence statement', 'v3 explanation statement')
LEVELS = {'CERTIFIED': 'C', 'ADVANCED': 'B', 'SHOWCASE': 'A'}
LEVEL_FOR_POSITION = {p: 'C' if p <= 5 else 'B' if p <= 8 else 'A' for p in range(1, 11)}
QUESTION_ID = re.compile(r'(\d{2})\.(\d{2})\.([CBA])')
LETTER = re.compile(r'^\s*([A-Z])\.\)')
EXTRACT_SCHEMA = 'tca-v3-source-rows-1'


class SourceError(Exception):
    pass


def _cell(value):
    if value is None: return None
    if isinstance(value, float) and value.is_integer(): return int(value)
    return value


def rows_from_workbook(path):
    """Raw cell values of the source sheet, with Excel row numbers. Opens read-only; never saves."""
    import openpyxl
    path = Path(path)
    data = path.read_bytes()
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        if SHEET not in workbook.sheetnames: raise SourceError(f'Sheet {SHEET} not found; found {workbook.sheetnames}.')
        sheet = workbook[SHEET]
        rows = []
        for number, values in enumerate(sheet.iter_rows(values_only=True), 1):
            values = [_cell(v) for v in list(values)[:11]] + [None] * max(0, 11 - len(values))
            if any(v not in (None, '') for v in values): rows.append([number, values])
    finally:
        workbook.close()
    return {'schema': EXTRACT_SCHEMA, 'filename': path.name, 'sha256': hashlib.sha256(data).hexdigest(),
            'sheet': SHEET, 'rows': rows}


def load_source(path):
    """Accept the original .xlsx or a JSON extract made by tools/extract_v3_source.py."""
    path = Path(path)
    if path.suffix.lower() == '.json':
        data = json.loads(path.read_text(encoding='utf-8'))
        if data.get('schema') != EXTRACT_SCHEMA: raise SourceError('Unknown extract format.')
        return data
    return rows_from_workbook(path)


def _text(lines):
    """Join a column's lines exactly; spacer rows become empty lines; trailing spacers are dropped."""
    values = ['' if v is None else str(v) for v in lines]
    while values and values[-1] == '': values.pop()
    while values and values[0] == '': values.pop(0)
    return '\n'.join(values)


def parse(source):
    """Group raw rows into source records. Returns (header, records, anomalies)."""
    rows = source['rows']
    anomalies = []
    header = {'release': None, 'columns': None, 'notes': []}
    records = []
    current = None
    for number, values in rows:
        first = values[0]
        if current is None and isinstance(first, str) and first.startswith('TCA v3'):
            header['release'] = first; continue
        if current is None and first == COLUMNS[0]:
            header['columns'] = [v for v in values]
            if tuple(values) != COLUMNS: anomalies.append({'row': number, 'kind': 'header', 'detail': 'Column headings differ from the expected A–K layout.'})
            continue
        if isinstance(first, int) and values[1]:
            current = {'master_order': first, 'row': number, 'last_row': number, 'values': values,
                       'I': [values[8]], 'J': [values[9]], 'K': [values[10]], 'rows': [[number, values]]}
            records.append(current)
            continue
        if current is None:
            header['notes'].append({'row': number, 'values': values})
            continue
        if any(v not in (None, '') for v in values[:8]):
            anomalies.append({'row': number, 'kind': 'orphan', 'detail': 'Continuation row has content in columns A–H.', 'master_order': current['master_order']})
        current['I'].append(values[8]); current['J'].append(values[9]); current['K'].append(values[10])
        current['last_row'] = number
        current['rows'].append([number, values])
    for record in records:
        v = record['values']
        record.update({'question_id': str(v[1]).strip(), 'order': v[2], 'section': v[3], 'level': v[4], 'language': v[5],
                       'groups': [g.strip() for g in str(v[6] or '').split(';') if g.strip()], 'product_line': v[7],
                       'requirement': _text(record['I']), 'evidence': _text(record['J']), 'explanation': _text(record['K'])})
        # Row-level gaps in I/J/K after the record row are kept as empty lines in the field text above.
    return header, records, anomalies


def clauses(requirement_text):
    """Derived presentation only: lettered sub-requirements (A.), B.) …) as evidence units.

    Unlettered lines after a lettered line stay with that line; text before the first letter is an introduction.
    Without letters the whole requirement is one unit. The exact text is never changed.
    """
    intro, items = [], []
    for line in requirement_text.split('\n'):
        match = LETTER.match(line)
        if match: items.append({'label': match.group(1), 'lines': [line]})
        elif items: items[-1]['lines'].append(line)
        else: intro.append(line)
    if not items: return '', [{'label': '', 'text': requirement_text}]
    return _text(intro), [{'label': i['label'], 'text': _text(i['lines'])} for i in items]


def reconcile(header, records, anomalies):
    """Counts and structural checks. Every problem is reported; an empty 'errors' list means the source is consistent."""
    errors = list(a['detail'] + f" (row {a['row']})" for a in anomalies if a['kind'] in ('orphan', 'header'))
    orders = [r['master_order'] for r in records]
    if orders != list(range(1, len(records) + 1)): errors.append('Master Order is not a continuous 1..n sequence.')
    ids = Counter(r['question_id'] for r in records)
    sections = sorted({r['section'] for r in records})
    groups = sorted({g for r in records for g in r['groups']})
    for r in records:
        match = QUESTION_ID.fullmatch(r['question_id'])
        if not match: errors.append(f"Master Order {r['master_order']}: QuestionID {r['question_id']!r} has an unexpected form."); continue
        if not isinstance(r['order'], int) or not 1 <= r['order'] <= 10: errors.append(f"Master Order {r['master_order']}: order {r['order']!r} outside 1–10.")
        elif int(match.group(2)) != r['order']: errors.append(f"Master Order {r['master_order']}: QuestionID position differs from order.")
        if not str(r['section']).startswith(match.group(1) + '.'): errors.append(f"Master Order {r['master_order']}: QuestionID section differs from section {r['section']!r}.")
        level = LEVELS.get(str(r['level']).upper())
        if level is None: errors.append(f"Master Order {r['master_order']}: unknown level {r['level']!r}.")
        elif level != match.group(3) or (isinstance(r['order'], int) and LEVEL_FOR_POSITION.get(r['order']) != level):
            errors.append(f"Master Order {r['master_order']}: level {r['level']} does not match position {r['order']}.")
        if not r['groups']: errors.append(f"Master Order {r['master_order']}: no TCA group.")
        if not r['requirement'].strip(): errors.append(f"Master Order {r['master_order']}: empty requirement text.")
    per_group = {}
    for group in groups:
        applicable = [r for r in records if group in r['groups']]
        counted = Counter(r['question_id'] for r in applicable)
        overlap = sorted(q for q, n in counted.items() if n > 1)
        missing = sorted(set(ids) - set(counted))
        positions = defaultdict(set)
        for r in applicable: positions[r['section']].add(r['order'])
        incomplete = sorted(s for s in sections if positions.get(s) != set(range(1, 11)))
        per_group[group] = {'records': len(applicable), 'questions': len(counted), 'overlap': overlap, 'missing': missing,
                            'incomplete_sections': incomplete}
        if overlap: errors.append(f'{group}: more than one variant for {", ".join(overlap)}.')
        if missing: errors.append(f'{group}: no variant for {", ".join(missing[:10])}{" …" if len(missing) > 10 else ""}.')
        if incomplete: errors.append(f'{group}: sections without exactly positions 1–10: {", ".join(incomplete)}.')
    variants = {q: [{'master_order': r['master_order'], 'row': r['row'], 'groups': r['groups']} for r in records if r['question_id'] == q]
                for q, n in sorted(ids.items()) if n > 1}
    multi_row = [r for r in records if r['last_row'] > r['row']]
    # Observations only (wording is never corrected): lettering that skips or does not start at A.
    lettering = []
    for r in records:
        labels = [u['label'] for u in clauses(r['requirement'])[1] if u['label']]
        expected = [chr(ord('A') + i) for i in range(len(labels))]
        if labels and labels != expected:
            lettering.append({'master_order': r['master_order'], 'question_id': r['question_id'], 'row': r['row'], 'labels': ''.join(labels)})
    return {
        'release': header.get('release'), 'notes': header.get('notes'),
        'records': len(records), 'questions': len(ids), 'sections': len(sections), 'groups': groups,
        'levels': dict(Counter(r['level'] for r in records)),
        'languages': dict(Counter(r['language'] for r in records)), 'product_lines': dict(Counter(r['product_line'] for r in records)),
        'variant_questions': variants, 'extra_variants': len(records) - len(ids),
        'per_group': per_group,
        'continuation': {'multi_row_records': len(multi_row),
                         'with_evidence_continuation': sum(any(v not in (None, '') for v in r['J'][1:]) for r in records),
                         'with_explanation_continuation': sum(any(v not in (None, '') for v in r['K'][1:]) for r in records),
                         'lettered_records': sum(bool(clauses(r['requirement'])[1][0]['label']) for r in records)},
        'lettering_observations': lettering,
        'anomalies': anomalies, 'errors': errors,
    }


def read_and_reconcile(path):
    source = load_source(path)
    header, records, anomalies = parse(source)
    return source, records, reconcile(header, records, anomalies)

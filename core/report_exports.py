"""Report workbooks and multi-event PDFs for Kaizen events and the sensei queue.

These are operational reports, not formula-compatible Global Kaizen workbooks. Team names and roles
are included; contact details, hours and benefits stay in the restricted archive.
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .a3 import RENDERER_VERSION, _photos, build_pdf_many, frozen_content, report_blocks
from .access import report_visible
from .kaizen import REVIEW_STATE_LABELS, gate_summary, review_state, team_display
from .measurements import compare_observations
from .methodology import rules_for

HEADER_FILL=PatternFill('solid',fgColor='0B1A48')

def _sheet(wb,title,headers,first=False):
    sheet=wb.active if first else wb.create_sheet(title)
    sheet.title=title;sheet.append(headers)
    for cell in sheet[1]: cell.font=Font(bold=True,color='FFFFFF');cell.fill=HEADER_FILL
    return sheet

def events_xlsx(events,user):
    wb=Workbook()
    overview=_sheet(wb,'Events',['Kaizen ID','Title','Site','Methodological type','Upload type (DL_UPLOAD)','Upload category (DL_UPLOAD)','Input view',
        'Workflow status','Review state','Gates met','Gates total','Submitted','Owner','Linked TCA requirements','Revision','Stored report versions'],first=True)
    summary=_sheet(wb,'A3 summary',['Kaizen ID','Block','Heading','Complete text'])
    actions=_sheet(wb,'Actions',['Kaizen ID','Phase','Action','Owner','Team member','Start','End','Progress %','Key action'])
    gates=_sheet(wb,'Gates',['Kaizen ID','Phase','Gate','Met'])
    team=_sheet(wb,'Team',['Kaizen ID','Name','Kaizen role'])
    results=_sheet(wb,'Results',['Kaizen ID','Measure','Population','Unit','Before average','Before points','After average','After points','Change'])
    review=_sheet(wb,'Review',['Kaizen ID','Phase','Type','Text','Sensei / author','Event revision','Resolved','Date'])
    for event in events:
        key=f'KZ-{event.pk}';rules=rules_for(event.profile);gate=gate_summary(event)
        stored=[f'{r.paper} v{r.version} (r{r.event_revision})' for r in event.a3_reports.order_by('version') if report_visible(r,user)]
        overview.append([key,event.title,event.site.name,rules['label'],rules['upload_type'],rules['upload_category'],event.get_input_profile_display(),
            event.status,REVIEW_STATE_LABELS[review_state(event)],gate['done'],gate['total'],event.submitted_at.isoformat() if event.submitted_at else '',
            event.data.get('owner',''),', '.join(event.linked_requirements.values_list('code',flat=True)),event.revision,', '.join(stored)])
        for number,heading,text in report_blocks(frozen_content(event)): summary.append([key,number,heading,text])
        for a in event.actions.select_related('assignee__user').order_by('start_date','pk'):
            actions.append([key,a.get_phase_display(),a.description,a.owner,team_display(a.assignee) if a.assignee else '',a.start_date,a.end_date,a.progress,'Yes' if a.key_action else 'No'])
        for g in gate['gates']: gates.append([key,g['phase'],g['label'],'Yes' if g['ok'] else 'No'])
        for p in event.participants.select_related('user'): team.append([key,team_display(p),p.kaizen_role])
        for c in compare_observations(event.observations.all()):
            results.append([key,c['measure'],c['population'],c['unit'],_num(c['before_average']),c['before_count'],_num(c['after_average']),c['after_count'],_num(c['difference'])])
        for n in event.review_notes.select_related('actor').order_by('created_at'):
            review.append([key,n.stage,n.get_kind_display(),n.text,n.actor.get_full_name() or n.actor.username,n.event_revision,'Yes' if n.resolved else 'No',n.created_at.isoformat()])
    for sheet in wb:
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if isinstance(cell.value,str): cell.data_type='s'  # user text is never an executable formula
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for column,width in zip('ABCDEFGHIJ',(12,40,24,24,40,16,16,16,16,16)): sheet.column_dimensions[column].width=width
    output=BytesIO();wb.save(output);output.seek(0)
    return output

def _num(value):
    return float(value) if value is not None else None

def events_pdf(events,user,paper='A3'):
    """Each event as its stored finalised version for the current revision, else a marked working copy."""
    items=[]
    for event in events:
        stored=event.a3_reports.filter(event_revision=event.revision,paper=paper,renderer_version=RENDERER_VERSION).order_by('-version').first()
        if stored and report_visible(stored,user): items.append((stored.snapshot,_photos(event),True))
        else: items.append((frozen_content(event),_photos(event),False))
    return build_pdf_many(items,paper,f'TCA Progress sensei review · {len(events)} events','TCA Progress · selected sensei review events · finalised versions or marked working copies')

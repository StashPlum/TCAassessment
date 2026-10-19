"""Stage-three journeys, frozen exports and site boundaries."""
from copy import deepcopy
from io import BytesIO
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook
from pypdf import PdfReader
from .assessment_documents import xlsx_bytes
from .models import (Assessment, AssessmentSession, Coverage, Evidence, Membership,
                     PreparationNeedVersion, Requirement, Response, SessionRevision, Site,
                     SiteQuestionVersion)


class StageThreeTests(TestCase):
    def test_export_treats_formula_like_source_text_as_text(self):
        sheet=load_workbook(BytesIO(xlsx_bytes('QA', [{'Assessor comment':'=2+2'}], 'one record')),
                            read_only=True,data_only=False).active
        self.assertEqual(sheet['M6'].value,'=2+2')
        self.assertEqual(sheet['M6'].data_type,'s')

    @classmethod
    def setUpTestData(cls):
        call_command('import_catalog', verbosity=0)
        cls.site=Site.objects.create(name='Synthetic Stage Three',code='stage-three')
        cls.other=Site.objects.create(name='Private Stage Three',code='private-three')
        User=get_user_model()
        cls.auditor=User.objects.create_user('three-auditor')
        cls.manager=User.objects.create_user('three-manager')
        cls.outsider=User.objects.create_user('three-outsider')
        Membership.objects.create(user=cls.auditor,site=cls.site,role='assessor')
        Membership.objects.create(user=cls.manager,site=cls.site,role='manager')
        Membership.objects.create(user=cls.outsider,site=cls.other,role='assessor')
        cls.req=Requirement.objects.get(pk='1.C.1')

    def _xlsx(self,response):
        self.assertEqual(response.status_code,200)
        self.assertEqual(response['Content-Type'],'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        return load_workbook(BytesIO(response.content),read_only=True,data_only=True).active

    def _pdf(self,response):
        self.assertEqual(response.status_code,200)
        self.assertEqual(response['Content-Type'],'application/pdf')
        return '\n'.join(page.extract_text() for page in PdfReader(BytesIO(response.content)).pages)

    def test_preparation_version_freeze_and_full_selected_export(self):
        url=reverse('preparation_need',args=[self.site.code,self.req.code])
        self.client.force_login(self.auditor)
        official=self.req.detail_text
        body={'evidence_items':'Shift checklist\nAttendance log','owner':'Site operations',
              'period_start':'2026-01-01','period_end':'2026-09-30',
              'source_note':'Local document register v4','change_reason':'Prepare annual review',
              'expected_version':'0'}
        self.assertEqual(self.client.post(url,body).status_code,302)
        first=PreparationNeedVersion.objects.get(site=self.site,requirement=self.req)
        self.assertEqual(first.diff['owner']['after'],'Site operations')
        self.assertEqual(self.client.post(url,body).status_code,409)
        assessment=Assessment.objects.create(site=self.site,version=1,created_by=self.auditor)
        frozen=deepcopy(assessment.preparation_snapshot['requirements'][self.req.code])
        body.update({'evidence_items':'Revised site checklist','expected_version':'1',
                     'change_reason':'Site procedure updated'})
        self.assertEqual(self.client.post(url,body).status_code,302)
        assessment.refresh_from_db()
        self.assertEqual(assessment.preparation_snapshot['requirements'][self.req.code],frozen)
        self.assertEqual(self.req.detail_text,official)
        first.owner='Tamper'
        with self.assertRaises(ValueError):first.save()
        self.client.post(reverse('preparation',args=[self.site.code]),
                         {'selection_action':'add','selected':[self.req.code],'page':'1'})
        last=Requirement.objects.order_by('category_id','position').last()
        self.client.post(reverse('preparation',args=[self.site.code]),
                         {'selection_action':'add','selected':[last.code],'page':'6'})
        xlsx=self._xlsx(self.client.get(reverse('preparation_export',args=[self.site.code,'xlsx']),
                                        {'scope':'selected'}))
        self.assertEqual(xlsx['B2'].value,'Selected requirements across pages (2)')
        self.assertEqual(xlsx['B3'].value,2)
        self.assertEqual({xlsx['F6'].value,xlsx['F7'].value},{self.req.code,last.code})
        self.assertEqual(xlsx['J6'].value,'Site operations')
        self.assertIn(official,xlsx['G6'].value)
        self.assertIn('2026-01-01',xlsx['K6'].value)
        pdf=self._pdf(self.client.get(reverse('preparation_export',args=[self.site.code,'pdf']),
                                      {'scope':'selected'}))
        for value in ('CEVA Logistics',self.req.code,last.code,'Site operations'):
            self.assertIn(value,pdf)
        self.client.force_login(self.outsider)
        for endpoint in (url,reverse('preparation_export',args=[self.site.code,'xlsx'])):
            self.assertEqual(self.client.get(endpoint).status_code,404)

    def test_plan_manager_alternative_rescheduling_and_started_snapshot(self):
        self.client.force_login(self.auditor)
        create=reverse('sessions',args=[self.site.code])
        body={'date':'2026-11-12','method':'On-site','purpose':'Annual review',
              'scope_text':'Safety section','participants':'Auditor and local team',
              'blocks_text':'09:00-09:45 | 2 | Safety and operations',
              'notes':'Please prepare records','owner':'Local coordinator',
              'requirements':[self.req.code]}
        self.assertEqual(self.client.post(create,body).status_code,302)
        plan=AssessmentSession.objects.get(site=self.site)
        self.assertEqual(plan.time_blocks[0]['people'],2)
        detail=reverse('session_detail',args=[self.site.code,plan.pk])
        self.assertEqual(self.client.post(detail,{'action':'review','availability':'confirmed',
            'manager_note':'Ready','revision':0}).status_code,404)
        self.client.force_login(self.manager)
        self.assertEqual(self.client.post(detail,{'action':'review','availability':'alternative',
            'alternate_date':'2026-11-13','backup_participants':'Backup supervisor',
            'manager_note':'Peak shift on proposed date','revision':0}).status_code,302)
        plan.refresh_from_db();self.assertEqual(plan.availability,'alternative')
        self.assertEqual(self.client.post(reverse('new_draft',args=[self.site.code]),
                                          {'session_id':plan.pk}).status_code,422)
        self.assertEqual(self.client.post(detail,{'action':'review','availability':'confirmed',
            'manager_note':'Backup confirmed','revision':1}).status_code,302)
        self.assertEqual(self.client.post(reverse('new_draft',args=[self.site.code]),
                                          {'session_id':plan.pk}).status_code,302)
        assessment=Assessment.objects.get(site=self.site,version=1)
        frozen=deepcopy(assessment.preparation_snapshot)
        self.client.force_login(self.auditor)
        changed={**body,'date':'2026-11-13','change_reason':'Late shift change',
                 'impact':'Supervisor unavailable at 09:00','followup_owner':'Local coordinator',
                 'revision':'2','action':'change'}
        self.assertEqual(self.client.post(detail,changed).status_code,302)
        self.assertEqual(self.client.post(detail,changed).status_code,409)
        plan.refresh_from_db();assessment.refresh_from_db()
        self.assertEqual(plan.availability,'proposed')
        self.assertEqual(assessment.preparation_snapshot,frozen)
        revision=SessionRevision.objects.get(session=plan,revision=3)
        self.assertEqual(revision.reason,'Late shift change')
        self.assertEqual(revision.followup_owner,'Local coordinator')
        self.assertEqual(revision.before['date'],'2026-11-12')
        self.assertEqual(revision.after['date'],'2026-11-13')
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(detail).status_code,404)
        self.assertEqual(self.client.post(detail,changed).status_code,404)

    def test_conduct_views_autosave_retest_and_history_frozen_export(self):
        evidence=Evidence.objects.create(site=self.site,title='Historical shift record',
            evidence_type='document',owner='Site operations',created_by=self.auditor)
        for clause in self.req.clauses.all():
            Coverage.objects.create(evidence=evidence,clause=clause,accepted=True,
                decision='Reviewed',reviewed_by=self.auditor,reviewed_at=timezone.now())
        first=Assessment.objects.create(site=self.site,version=1,created_by=self.auditor)
        Response.objects.create(assessment=first,requirement=self.req,code='A',
            conclusion='Historical partial control',advisory_follow_up='Review revised checklist',
            advisory_retest_status='due_next',updated_by=self.auditor)
        self.client.force_login(self.manager)
        self.client.post(reverse('approve',args=[self.site.code]),{'rationale':'Dated local record'})
        first.refresh_from_db();historical=deepcopy(first.snapshot)
        SiteQuestionVersion.objects.create(site=self.site,requirement=self.req,version=1,
            kind='local_guidance',question_text='Revised question for new assessment only',
            evidence_items=['New local document'],change_reason='New site process',
            author=self.auditor)
        self.client.force_login(self.auditor)
        self.client.post(reverse('new_draft',args=[self.site.code]),{'rationale':'Retest'})
        second=Assessment.objects.get(site=self.site,version=2)
        list_url=reverse('assessment',args=[self.site.code])
        for view,count in [('all',20),('section',10),('single',1)]:
            page=self.client.get(list_url,{'assessment':second.pk,'view':view})
            self.assertEqual(page.status_code,200)
            self.assertEqual(len(page.context['page'].object_list),count)
            self.assertEqual(page.context['assessment'].pk,second.pk)
        self.assertEqual(self.client.get(list_url,{'assessment':second.pk,'evidence_type':'document'}).context['page'].paginator.count,1)
        detail=reverse('respond',args=[self.site.code,self.req.code])
        self.assertContains(self.client.get(detail,{'assessment':second.pk}),'Review revised checklist')
        failed=self.client.post(detail,{'assessment_id':second.pk,'autosave':'1','code':'Y',
            'conclusion':'New check','revision':'0'})
        self.assertEqual(failed.status_code,422)
        self.assertFalse(failed.json()['saved'])
        saved=self.client.post(detail,{'assessment_id':second.pk,'autosave':'1','code':'Y',
            'conclusion':'New check','retest_confirm':'on','revision':'0'})
        self.assertEqual(saved.status_code,200)
        self.assertTrue(saved.json()['saved'])
        self.assertEqual(saved.json()['revision'],1)
        self.assertEqual(self.client.post(detail,{'assessment_id':second.pk,'autosave':'1',
            'code':'N','revision':'0'}).status_code,409)
        for view in ('all','section','single'):
            page=self.client.get(list_url,{'assessment':second.pk,'view':view,'answer':'Y'})
            self.assertEqual(page.context['page'].paginator.count,1)
            self.assertEqual(page.context['page'].object_list[0].response_code,'Y')
        self.assertContains(self.client.get(detail,{'assessment':second.pk}),'New check')
        history=reverse('assessment_history',args=[self.site.code])
        page=self.client.get(history,{'year':str(first.created_at.year),'site':self.site.code,
                                      'version':'1','auditor':self.auditor.username})
        self.assertEqual(page.context['page'].paginator.count,1)
        export=reverse('assessment_history_export',args=[self.site.code,'xlsx'])
        sheet=self._xlsx(self.client.post(export,{'assessment_ids':[first.pk,second.pk]}))
        self.assertEqual(sheet['B3'].value,300)
        self.assertEqual(sheet['L6'].value,'A')
        self.assertNotEqual(sheet['H6'].value,'Revised question for new assessment only')
        self.assertEqual(sheet['H156'].value,'Revised question for new assessment only')
        self.assertEqual(sheet['M6'].value,'Historical partial control')
        self.assertIn('Historical shift record',sheet['P6'].value)
        self.assertEqual(sheet['D6'].value,'advisory-v2')
        text=self._pdf(self.client.post(reverse('assessment_history_export',args=[self.site.code,'pdf']),
                                        {'assessment_ids':[first.pk]}))
        self.assertIn('Historical partial control',text)
        self.assertIn('Review revised checklist',text)
        self.assertEqual(self.client.post(detail,{'assessment_id':first.pk,'code':'N',
            'revision':'0'}).status_code,403)
        first.refresh_from_db();self.assertEqual(first.snapshot,historical)
        foreign=Assessment.objects.create(site=self.other,version=1,created_by=self.outsider)
        self.assertEqual(self.client.post(export,{'assessment_ids':[first.pk,foreign.pk]}).status_code,404)
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(list_url,{'assessment':second.pk}).status_code,404)
        self.assertEqual(self.client.get(detail,{'assessment':second.pk}).status_code,404)
        self.assertEqual(self.client.get(history).status_code,404)

    def test_restricted_evidence_reference_is_redacted_for_team_export(self):
        staff=get_user_model().objects.create_user('three-team')
        Membership.objects.create(user=staff,site=self.site,role='staff')
        evidence=Evidence.objects.create(site=self.site,title='Restricted HR interview',
            evidence_type='interview',classification='hr',owner='HR',created_by=self.auditor)
        Coverage.objects.create(evidence=evidence,clause=self.req.clauses.first(),accepted=True,
            decision='Reviewed',reviewed_by=self.auditor,reviewed_at=timezone.now())
        assessment=Assessment.objects.create(site=self.site,version=1,created_by=self.auditor)
        Response.objects.create(assessment=assessment,requirement=self.req,code='N',updated_by=self.auditor)
        self.client.force_login(self.manager)
        self.client.post(reverse('approve',args=[self.site.code]),{'rationale':'Dated record'})
        url=reverse('assessment_history_export',args=[self.site.code,'xlsx'])
        self.client.force_login(staff)
        sheet=self._xlsx(self.client.post(url,{'assessment_ids':[assessment.pk]}))
        self.assertNotIn('Restricted HR interview',sheet['P6'].value)
        detail=self.client.get(reverse('respond',args=[self.site.code,self.req.code]),
                               {'assessment':assessment.pk})
        self.assertNotContains(detail,'Restricted HR interview')
        self.client.force_login(self.auditor)
        sheet=self._xlsx(self.client.post(url,{'assessment_ids':[assessment.pk]}))
        self.assertIn('Restricted HR interview',sheet['P6'].value)

import json
import re
from pathlib import Path
from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from .models import Assessment, Category, Clause, Evidence, Kaizen, Membership, OtherAudit, Requirement, Response, Site
from .scoring import category_result
from .quality import results_criterion, sustain_types_criterion
from .models import Coverage, Observation, Benefit
from decimal import Decimal
from datetime import date, timedelta
from openpyxl import load_workbook
from io import BytesIO
from .source import star_criteria, kaizen_record
from .forms import KaizenStageForm, SessionForm

class SourceTests(TestCase):
    def test_exact_import_and_repeat(self):
        call_command('import_catalog',verbosity=0)
        call_command('import_catalog',verbosity=0)
        data=json.loads((Path(__file__).resolve().parents[1]/'sources/extracted/TCA assessment work categorisation.json').read_text(encoding='utf-8'))
        self.assertEqual(Category.objects.count(),15)
        self.assertEqual(Requirement.objects.count(),150)
        self.assertEqual(Clause.objects.count(),494)
        for row in data['requirements']:
            req=Requirement.objects.get(pk=row['id'])
            self.assertEqual(req.detail_text,row['detail_text'])
            self.assertEqual(req.reasoning,row['reasoning'] or '')
            self.assertEqual(req.expected_evidence,row['expected_evidence'] or '')
            self.assertEqual(req.remote_method,row['remote_method'] or '')
            self.assertEqual([(c.source_address,c.exact_text) for c in req.clauses.all()],[(c['address'],c['text']) for c in row['clauses']])
        assigned={c['address'].split('!')[1] for row in data['requirements'] for c in row['clauses']}
        headings={c['source'].split('!')[1].split(':')[1] for c in data['categories']}
        unassigned=[cell['address'] for cell in data['sheets'][0]['cells'] if cell['address'].startswith('B')
                    and isinstance(cell['value'],str) and cell['value'].strip()
                    and cell['address'] not in assigned|headings|{'B3'}]
        self.assertEqual(unassigned,[])

    def test_kaizen_source_metadata(self):
        source=kaizen_record()
        self.assertEqual(len(source['sheets']),15)
        self.assertEqual(len(star_criteria()),10)
        formulas=[c['value'] for s in source['sheets'] for c in s['cells'] if c['type']=='f']
        self.assertTrue(any(isinstance(f,dict) and f.get('kind')=='ArrayFormula' and f.get('text') for f in formulas))

class ScoringTests(TestCase):
    def test_boundaries_and_unresolved_a(self):
        self.assertIsNone(category_result({i:'Y' for i in range(1,5)})['earned'])
        self.assertEqual(category_result({i:'Y' for i in range(1,6)})['earned'],'Certified')
        self.assertEqual(category_result({**{i:'Y' for i in range(1,6)},6:'A',7:'Y',8:'Y'})['earned'],'Certified')
        self.assertEqual(category_result({i:'Y' for i in range(1,9)})['earned'],'Advanced')
        self.assertEqual(category_result({**{i:'Y' for i in range(1,9)},9:'N',10:'Y'})['earned'],'Advanced')
        self.assertEqual(category_result({i:'Y' for i in range(1,11)})['earned'],'Showcase')

    def test_star_boundaries(self):
        class Point:
            def __init__(self,phase,day):
                self.phase=phase;self.measure='checks';self.unit='count';self.population='shifts'
                self.period_start=date(2026,1,1)+timedelta(days=day)
                self.period_end=self.period_start;self.value=Decimal('10')
        data={'target':'9','direction':'up'}
        before=[Point('before',0)]
        self.assertFalse(results_criterion(data,before+[Point('after',i*6) for i in range(6)])) # exactly 30 days
        self.assertFalse(results_criterion(data,before+[Point('after',i*8) for i in range(5)])) # only 5 points
        self.assertTrue(results_criterion(data,before+[Point('after',i*7) for i in range(6)]))
        self.assertFalse(sustain_types_criterion({'document_types':'work instruction, WORK INSTRUCTION'}))
        self.assertTrue(sustain_types_criterion({'document_types':'work instruction, performance board'}))

class JourneyTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        call_command('import_catalog',verbosity=0)
        cls.site=Site.objects.create(name='Fictional Test Site',code='fictional-test')
        cls.other=Site.objects.create(name='Other Fictional Site',code='other-fictional')
        cls.user=get_user_model().objects.create_user('tester',password='some-test-password')
        cls.outsider=get_user_model().objects.create_user('outsider',password='some-test-password')
        for role in ('staff','assessor','coach','manager'):
            Membership.objects.create(user=cls.user,site=cls.site,role=role)
        Membership.objects.create(user=cls.outsider,site=cls.other,role='staff')

    def test_journey_history_and_cross_site(self):
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse('assessment',args=[self.site.code])),'Start assessment')
        self.client.post(reverse('new_draft',args=[self.site.code]),{})
        self.assertEqual(self.client.get(reverse('assessment',args=[self.site.code])).status_code,200)
        req=Requirement.objects.get(pk='1.C.1')
        obj=Assessment.objects.get(site=self.site)
        save=reverse('respond',args=[self.site.code,req.code])
        self.assertEqual(self.client.post(save,{'code':'N','conclusion':'Synthetic gap','revision':0}).status_code,302)
        self.assertEqual(Response.objects.get(assessment=obj,requirement=req).code,'N')
        # A stale form cannot replace the saved conclusion.
        self.client.post(save,{'code':'Y','revision':0})
        self.assertEqual(Response.objects.get(assessment=obj,requirement=req).code,'N')
        self.client.post(reverse('approve',args=[self.site.code]),{'rationale':'Fictional dated baseline','confirm_partial':'on'})
        obj.refresh_from_db();self.assertEqual(obj.status,'approved')
        self.client.get(reverse('assessment',args=[self.site.code]))
        self.assertEqual(Assessment.objects.filter(site=self.site).count(),1)
        self.client.post(reverse('new_draft',args=[self.site.code]),{'rationale':'Reassessment after Kaizen'})
        new=Assessment.objects.get(site=self.site,version=2)
        self.assertFalse(Response.objects.filter(assessment=new).exists())
        self.client.post(save,{'code':'N','revision':0,'conclusion':'Still open'})
        self.client.post(save,{'code':'Y','revision':0,'conclusion':'Stale form'})
        self.assertEqual(Response.objects.get(assessment=new,requirement=req).code,'N')
        evidence=Evidence.objects.create(site=self.site,title='Fictional accepted observations',evidence_type='observation',owner='Demo',created_by=self.user)
        for clause in req.clauses.all(): Coverage.objects.create(evidence=evidence,clause=clause,accepted=True,reviewed_by=self.user,reviewed_at=timezone.now())
        self.client.post(save,{'code':'Y','revision':1,'conclusion':'Synthetic verified change'})
        self.assertEqual(Response.objects.get(assessment=new,requirement=req).code,'Y')
        self.assertEqual(Response.objects.get(assessment=obj,requirement=req).code,'N')
        self.client.post(reverse('approve',args=[self.site.code]),{'rationale':'Fictional reassessment','confirm_partial':'on'})
        new.refresh_from_db();self.assertEqual(new.snapshot['accepted_codes']['1.C.1'],'Y')
        self.assertEqual(len(new.snapshot['coverage']),req.clauses.count())
        self.assertEqual(new.snapshot['responses'][0]['conclusion'],'Synthetic verified change')
        Coverage.objects.filter(evidence=evidence).update(accepted=False)
        new.refresh_from_db();self.assertEqual(new.snapshot['accepted_codes']['1.C.1'],'Y')
        self.assertEqual(len(new.snapshot['coverage']),req.clauses.count())
        self.assertIn('1.C.1',self.client.get(reverse('report',args=[self.site.code])).json()['unverified_yes'])
        self.assertEqual(self.client.get(reverse('site',args=[self.other.code])).status_code,404)
        self.assertEqual(self.client.get(reverse('report',args=[self.other.code])).status_code,404)
        self.assertEqual(self.client.get(reverse('assessment_archive',args=[self.other.code])).status_code,404)
        self.assertEqual(self.client.get(reverse('evidence',args=[self.other.code])).status_code,404)
        archive=self.client.get(reverse('assessment_archive',args=[self.site.code])).json()
        self.assertEqual(len(archive['requirements']),150)
        self.assertEqual(archive['requirements'][0]['clauses'][0]['exact_text'],req.clauses.first().exact_text)

    def test_kaizen_does_not_change_assessment(self):
        self.client.force_login(self.user)
        a=Assessment.objects.create(site=self.site,version=1,created_by=self.user)
        r=Requirement.objects.get(pk='1.C.1')
        Response.objects.create(assessment=a,requirement=r,code='N',updated_by=self.user)
        k=Kaizen.objects.create(site=self.site,title='Synthetic improvement',profile='A4',created_by=self.user)
        k.linked_requirements.add(r)
        url=reverse('kaizen_detail',args=[self.site.code,k.pk])
        self.client.post(url,{'revision':0,'action':'verify','overall_problem':'Synthetic'})
        k.refresh_from_db();self.assertEqual(k.status,'proposed')
        self.assertEqual(Response.objects.get(assessment=a,requirement=r).code,'N')
        self.assertEqual(self.client.get(reverse('kaizen_archive',args=[self.other.code,k.pk])).status_code,404)

    def test_evidence_coverage_and_financial_scope(self):
        self.client.force_login(self.user)
        req=Requirement.objects.get(pk='1.C.1')
        item=Evidence.objects.create(site=self.site,title='Fictional audit',evidence_type='audit',owner='Demo',created_by=self.user)
        link=Coverage.objects.create(evidence=item,clause=req.clauses.first())
        self.assertFalse(link.accepted)
        self.client.post(reverse('review_coverage',args=[self.site.code,link.pk]),{'decision':'accept','reason':'Named clause checked'})
        link.refresh_from_db();self.assertFalse(link.accepted)
        OtherAudit.objects.create(evidence=item,scope='Synthetic site audit',period_start=date(2026,1,1),period_end=date(2026,1,31),tested_control='1.C.1',finding='No exception')
        self.client.post(reverse('review_coverage',args=[self.site.code,link.pk]),{'decision':'accept','reason':'Named clause checked'})
        link.refresh_from_db();self.assertTrue(link.accepted)
        self.assertFalse(Response.objects.filter(requirement=req).exists())
        k=Kaizen.objects.create(site=self.site,title='Synthetic',profile='A3',created_by=self.user)
        Benefit.objects.create(event=k,kind='savings',amount=Decimal('100'),unit='USD')
        self.client.force_login(self.outsider)
        self.assertEqual(self.client.get(reverse('download',args=[self.site.code,item.pk])).status_code,404)
        self.assertEqual(self.client.get(reverse('kaizen_archive',args=[self.site.code,k.pk])).status_code,404)
        self.client.force_login(self.user)
        archive=self.client.get(reverse('kaizen_archive',args=[self.site.code,k.pk])).json()
        self.assertEqual(len(archive['benefits']),1)

    def test_remote_session_restriction(self):
        form=SessionForm(data={'date':'2026-10-04','method':'remote video','requirements':['1.B.6']})
        self.assertFalse(form.is_valid())
        self.assertIn('remote assessment',str(form.errors).lower())

    def test_staff_cannot_see_benefit_details(self):
        staff=get_user_model().objects.create_user('staff_only',password='test-pass')
        Membership.objects.create(user=staff,site=self.site,role='staff')
        k=Kaizen.objects.create(site=self.site,title='Synthetic',profile='A4',created_by=self.user)
        Benefit.objects.create(event=k,kind='savings',amount=Decimal('12345'),unit='USD')
        self.client.force_login(staff)
        page=self.client.get(reverse('kaizen_detail',args=[self.site.code,k.pk]))
        self.assertNotContains(page,'12345')
        self.assertEqual(self.client.get(reverse('kaizen_archive',args=[self.site.code,k.pk])).json()['benefits'],'Restricted')
        self.assertEqual(self.client.post(reverse('kaizen_add',args=[self.site.code,k.pk,'benefit']),{}).status_code,404)

    def test_excel_archive_and_external_limit(self):
        self.client.force_login(self.user)
        k=Kaizen.objects.create(site=self.site,title='Fictional report',profile='A4',created_by=self.user,
            data={'overall_problem':'Long but complete fictional explanation','measure':'PPE checks'})
        url=reverse('kaizen_excel',args=[self.site.code,k.pk])
        response=self.client.get(url)
        self.assertEqual(response.status_code,200)
        wb=load_workbook(BytesIO(response.content))
        self.assertEqual(wb['PDCA Fields']['B5'].value,'Long but complete fictional explanation')
        self.assertIn('Compatibility',wb.sheetnames)
        k.title='X'*76;k.save()
        bad=self.client.get(url)
        self.assertEqual(bad.status_code,422)
        self.assertIn('75',str(bad.json()))

    def test_reassessment_acceptance_requires_approved_yes(self):
        self.client.force_login(self.user)
        req=Requirement.objects.get(pk='1.C.1')
        k=Kaizen.objects.create(site=self.site,title='Synthetic review',profile='A4',status='reassessment_pending',created_by=self.user)
        k.linked_requirements.add(req)
        url=reverse('kaizen_detail',args=[self.site.code,k.pk])
        self.client.post(url,{'revision':0,'action':'accept_reassessment'})
        k.refresh_from_db();self.assertEqual(k.status,'reassessment_pending')
        a=Assessment.objects.create(site=self.site,version=1,status='approved',created_by=self.user,approved_by=self.user)
        Response.objects.create(assessment=a,requirement=req,code='Y',updated_by=self.user)
        evidence=Evidence.objects.create(site=self.site,title='Synthetic accepted clauses',evidence_type='observation',owner='Demo',created_by=self.user)
        for clause in req.clauses.all(): Coverage.objects.create(evidence=evidence,clause=clause,accepted=True,reviewed_by=self.user,reviewed_at=timezone.now())
        self.client.post(url,{'revision':0,'action':'accept_reassessment'})
        k.refresh_from_db();self.assertEqual(k.status,'reassessment_accepted')

    def test_authorized_navigation_switching_and_admin_boundary(self):
        self.site.country_code='NL';self.site.grouping='Benelux';self.site.save()
        self.other.country_code='GB';self.other.save()
        self.client.force_login(self.user)
        page=self.client.get(reverse('home'))
        self.assertContains(page,'Netherlands')
        self.assertContains(page,'CAT confirmation pending')
        self.assertNotContains(page,'Other Fictional Site')
        self.assertNotContains(page,'United Kingdom')
        self.assertNotContains(page,'href="/administration/"')
        self.assertEqual(self.client.get(reverse('administration')).status_code,404)
        self.assertEqual(self.client.post(reverse('switch_site'),{'site_code':self.other.code}).status_code,404)
        self.assertRedirects(self.client.post(reverse('switch_site'),{'site_code':self.site.code,'area':'assessment'}),reverse('assessment',args=[self.site.code]))
        self.assertEqual(self.client.session['selected_site'],self.site.code)
        self.assertRedirects(self.client.post(reverse('switch_site'),{'site_code':self.site.code,'area':'https://example.com'}),reverse('home'))
        Membership.objects.create(user=self.user,site=self.site,role='admin')
        self.assertContains(self.client.get(reverse('administration')),'does not confer program approval authority')
        self.assertEqual(self.client.get(reverse('resources',args=[self.other.code])).status_code,404)

    def test_progressive_pdca_fields_survive_save_and_next(self):
        self.client.force_login(self.user)
        event=Kaizen.objects.create(site=self.site,title='Synthetic UX record',profile='A4',created_by=self.user,
                                   data={'overall_problem':'Keep original problem','coach_11':'Retain supplementary detail'})
        url=reverse('kaizen_detail',args=[self.site.code,event.pk])
        page=self.client.get(url)
        ids=re.findall(r'\bid="([^"]+)"',page.content.decode())
        self.assertEqual(len(ids),len(set(ids)))
        for key in KaizenStageForm().fields:
            self.assertContains(page,f'name="{key}"')
        data={key:page.context['form'][key].value() if page.context['form'][key].value() is not None else '' for key in page.context['form'].fields}
        data.update({'solution':'Synthetic countermeasure','action':'save'})
        self.assertEqual(self.client.post(url,data).status_code,302)
        event.refresh_from_db()
        self.assertEqual(event.data['overall_problem'],'Keep original problem')
        self.assertEqual(event.data['coach_11'],'Retain supplementary detail')
        invalid=self.client.post(reverse('kaizen_add',args=[self.site.code,event.pk,'observation']),
            {'phase':'after','measure':'Retained test measure','value':'0','unit':'count','population':'shifts',
             'period_start':'2026-10-12','period_end':'2026-10-01','data_source':'Synthetic'})
        self.assertEqual(invalid.status_code,200)
        self.assertContains(invalid,'Retained test measure')
        self.assertContains(invalid,'Period end must be on or after period start.')
        self.assertTrue(invalid.context['observation_form'].errors)
        self.assertEqual(event.observations.count(),0)
        req=Requirement.objects.get(pk='1.C.1')
        assessment=Assessment.objects.create(site=self.site,version=1,created_by=self.user)
        response=self.client.post(reverse('respond',args=[self.site.code,req.code]),{'revision':0,'code':'N','conclusion':'Synthetic','go':'next'})
        self.assertRedirects(response,reverse('respond',args=[self.site.code,'1.C.2']))
        self.assertEqual(Response.objects.get(assessment=assessment,requirement=req).code,'N')

    def test_empty_access_and_assessment_search_states(self):
        no_sites=get_user_model().objects.create_user('no_sites')
        self.client.force_login(no_sites)
        self.assertContains(self.client.get(reverse('home')),'No site access assigned')
        self.assertNotContains(self.client.get(reverse('home')),self.site.name)
        self.client.force_login(self.user)
        self.assertContains(self.client.get(reverse('assessment',args=[self.site.code]),{'q':'NONEXISTENT_TEST_REQUIREMENT'}),'No matching requirements')

    def test_staff_can_read_blocker_without_editing_or_restricted_evidence(self):
        staff=get_user_model().objects.create_user('read_only_staff')
        Membership.objects.create(user=staff,site=self.site,role='staff')
        req=Requirement.objects.get(pk='1.C.1')
        evidence=Evidence.objects.create(site=self.site,title='Restricted HR title',evidence_type='document',classification='hr',owner='Demo',created_by=self.user)
        Coverage.objects.create(evidence=evidence,clause=req.clauses.first())
        self.client.force_login(staff)
        url=reverse('respond',args=[self.site.code,req.code])
        self.assertContains(self.client.get(url),'Read-only access')
        self.assertNotContains(self.client.get(url),'Restricted HR title')
        self.assertEqual(self.client.post(url,{'code':'N','conclusion':'Blocked','revision':0}).status_code,404)
        self.assertFalse(Response.objects.filter(requirement=req).exists())

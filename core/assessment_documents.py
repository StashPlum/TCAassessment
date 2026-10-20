"""Readable operational assessment exports from frozen records; no CAT standing implied."""
from html import escape
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import CondPageBreak, Paragraph, SimpleDocTemplate, Spacer
from .assessment_config import preparation_for_assessment, question_for_assessment
from .models import Coverage, Evidence, Requirement

HEADERS=['Site','Assessment version','Record status','Rule version','Category','Requirement ID',
         'Official obligation','Frozen site question','Evidence needed','Document owner','Relevant period',
         'Answer','Assessor comment','Advisory follow-up','Retest status','Evidence references','Source address']


def _period(item):
    start=item.get('period_start') or ''
    end=item.get('period_end') or ''
    return f'{start} to {end}' if start or end else 'Period not specified'


def preparation_records(site, items):
    records=[]
    for item in items:
        req=item['requirement'];need=item['need'];question=item['question']
        records.append({'Site':site.name,'Assessment version':'Preparation, current site guidance',
            'Record status':'Working list','Rule version':'No response scored','Category':req.section_name,
            'Requirement ID':req.label,'Official obligation':req.detail_text,
            'Frozen site question':question['question_text'],
            'Evidence needed':'\n'.join(need['evidence_items']),
            'Document owner':need['owner'] or 'Not assigned','Relevant period':_period(need),
            'Answer':item['response'] or 'Unanswered','Assessor comment':'','Advisory follow-up':'',
            'Retest status':'','Evidence references':f"{item['accepted']} current accepted records",
            'Source address':f"{req.source_address}; preparation v{need['version']}; {need['source_note']}"})
    return records


def assessment_records(assessment, allowed_classifications=None):
    """Approved exports read response/evidence values from the dated snapshot."""
    scope=assessment.scope_snapshot.get('requirements') or list(Requirement.objects.values_list('code',flat=True))
    requirements=list(Requirement.objects.filter(code__in=scope).select_related('category').prefetch_related('clauses'))
    snap=assessment.snapshot if assessment.status=='approved' else {}
    responses={r['requirement']:r for r in snap.get('responses',[])} if assessment.status=='approved' else {
        r.requirement_id:{'code':r.code,'conclusion':r.conclusion,'advisory_follow_up':r.advisory_follow_up,
                          'advisory_retest_status':r.advisory_retest_status} for r in assessment.responses.all()}
    coverage_by_clause={}
    visible_ids=None
    if allowed_classifications is not None:
        visible_ids=set(Evidence.objects.filter(site=assessment.site,
            classification__in=allowed_classifications).values_list('pk',flat=True))
    if assessment.status=='approved':
        for link in snap.get('evidence_references',snap.get('coverage',[])):
            if visible_ids is not None and link.get('evidence_id') not in visible_ids:continue
            coverage_by_clause.setdefault(link['clause_id'],[]).append(
                f"{link['evidence_title']} v{link['evidence_version'] or '?'} ({link['source_address']}; "
                f"{'accepted' if link.get('accepted',True) else 'pending'}; {link['decision']})")
    else:
        # Only mutable drafts may consult live evidence. Dated approvals use retained rows alone.
        fallback=Coverage.objects.filter(evidence__site=assessment.site,accepted=True)
        if allowed_classifications is not None:fallback=fallback.filter(evidence__classification__in=allowed_classifications)
        for link in fallback.select_related('evidence','clause'):
            coverage_by_clause.setdefault(link.clause_id,[]).append(
                f"{link.evidence.title} v{link.evidence.version or '?'} (current/legacy reference; {link.clause.source_address})")
    records=[]
    for req in requirements:
        response=responses.get(req.code,{})
        question=question_for_assessment(assessment,req)
        need=preparation_for_assessment(assessment,req)
        refs=[]
        for clause in req.clauses.all(): refs.extend(coverage_by_clause.get(clause.pk,[]))
        records.append({'Site':assessment.site.name,'Assessment version':str(assessment.version),
            'Record status':assessment.status,'Rule version':assessment.policy_version,'Category':req.section_name,
            'Requirement ID':req.label,'Official obligation':req.detail_text,
            'Frozen site question':question['question_text'],'Evidence needed':'\n'.join(need['evidence_items']),
            'Document owner':need['owner'] or 'Not assigned','Relevant period':_period(need),
            'Answer':response.get('code') or 'Unanswered','Assessor comment':response.get('conclusion',''),
            'Advisory follow-up':response.get('advisory_follow_up',''),
            'Retest status':response.get('advisory_retest_status',''),
            'Evidence references':'\n'.join(dict.fromkeys(refs)) or 'No visible evidence reference in this snapshot',
            'Source address':f"{req.source_address}; {question['source_hash']}; site question v{question['version_id'] or 0}"})
    return records


def xlsx_bytes(title, records, scope_label):
    workbook=Workbook();sheet=workbook.active;sheet.title='Assessment records'
    sheet.append([title]);sheet.append(['Export scope',scope_label]);sheet.append(['Records',len(records)]);sheet.append([])
    sheet.append(HEADERS)
    for record in records:
        sheet.append([str(record.get(key,'') or '') for key in HEADERS])
        for cell in sheet[sheet.max_row]:
            if cell.data_type=='f':cell.data_type='s'  # A source/comment is text, never an executable workbook formula.
    navy=PatternFill('solid',fgColor='0B1D3B');light=PatternFill('solid',fgColor='EAF0F8')
    sheet['A1'].font=Font(bold=True,size=15,color='0B1D3B')
    for cell in sheet[5]: cell.fill=navy;cell.font=Font(bold=True,color='FFFFFF');cell.alignment=Alignment(wrap_text=True,vertical='top')
    widths={'A':24,'B':18,'C':17,'D':18,'E':28,'F':17,'G':65,'H':56,'I':46,'J':22,
            'K':26,'L':14,'M':44,'N':42,'O':19,'P':58,'Q':53}
    for col,width in widths.items():sheet.column_dimensions[col].width=width
    for row in sheet.iter_rows(min_row=6):
        for cell in row:cell.alignment=Alignment(wrap_text=True,vertical='top')
        if row[0].row%2==0:
            for cell in row:cell.fill=light
    sheet.freeze_panes='G6';sheet.auto_filter.ref=f'A5:Q{max(5,sheet.max_row)}'
    sheet.sheet_view.showGridLines=False
    output=BytesIO();workbook.save(output);return output.getvalue()


def pdf_bytes(title, records, scope_label):
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='TcaTitle',parent=styles['Title'],fontName='Helvetica-Bold',
        fontSize=17,leading=21,textColor=colors.HexColor('#0B1D3B'),spaceAfter=8))
    styles.add(ParagraphStyle(name='TcaRecord',parent=styles['Heading2'],fontName='Helvetica-Bold',
        fontSize=11,leading=14,textColor=colors.HexColor('#2458A6'),spaceBefore=14,spaceAfter=5))
    styles.add(ParagraphStyle(name='TcaBody',parent=styles['BodyText'],fontName='Helvetica',
        fontSize=8.5,leading=12,spaceAfter=5,splitLongWords=True))
    output=BytesIO();doc=SimpleDocTemplate(output,pagesize=A4,rightMargin=18*mm,leftMargin=18*mm,
        topMargin=22*mm,bottomMargin=18*mm,title=title,author='TCA Progress local pilot')
    story=[Paragraph(escape(title),styles['TcaTitle']),
           Paragraph(escape(f'Export scope: {scope_label}. {len(records)} requirements. Local operational record; CAT standing is separate.'),styles['TcaBody'])]
    keys=[('Official obligation','Official obligation'),('Frozen site question','Site question'),
          ('Evidence needed','Evidence needed'),('Document owner','Document owner'),('Relevant period','Relevant period'),
          ('Answer','Answer'),('Assessor comment','Assessor comment'),('Advisory follow-up','Advisory follow-up'),
          ('Retest status','Retest status'),('Evidence references','Evidence references'),('Source address','Source')]
    for record in records:
        story.append(CondPageBreak(65*mm))
        heading=f"{record['Requirement ID']} · {record['Category']} · {record['Record status']} · {record['Rule version']}"
        story.append(Paragraph(escape(heading),styles['TcaRecord']))
        for key,label in keys:
            # Escape before adding line breaks; source text is never markup.
            value=escape(str(record.get(key,'') or '—')).replace('\n','<br/>')
            story.append(Paragraph(f'<b>{escape(label)}:</b> {value}',styles['TcaBody']))
        story.append(Spacer(1,5*mm))
    def page(canvas,document):
        canvas.saveState();canvas.setFillColor(colors.HexColor('#0B1D3B'))
        canvas.setFont('Helvetica-Bold',9);canvas.drawString(18*mm, A4[1]-13*mm,'CEVA Logistics  |  TCA Progress')
        canvas.setStrokeColor(colors.HexColor('#CBD4E0'));canvas.line(18*mm,A4[1]-16*mm,A4[0]-18*mm,A4[1]-16*mm)
        canvas.setFont('Helvetica',8);canvas.drawRightString(A4[0]-18*mm,10*mm,f'Page {document.page}')
        canvas.restoreState()
    doc.build(story,onFirstPage=page,onLaterPages=page)
    return output.getvalue()
